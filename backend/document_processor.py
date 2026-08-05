from docx import Document


def extract_text_from_docx(file_path):

    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(
                paragraph.text
            )

    return "\n".join(paragraphs)


def create_docx(translated_text, output_path):

    document = Document()

    for paragraph in translated_text.split("\n"):
        document.add_paragraph(
            paragraph
        )

    document.save(output_path)
